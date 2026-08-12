"""Excel (.xlsx) exporter for DICOM inventory."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Sequence

from neuro_pipeline import __version__
from neuro_pipeline.inventory.models import (
    PROTOCOL_COLUMNS,
    SERIES_INVENTORY_COLUMNS,
    InventoryResult,
)

HEADER_FILL = "1F7A3A"
HEADER_FONT_COLOR = "FFFFFF"
ALT_ROW_FILL = "F3F7F4"


def export_inventory_xlsx(
    result: InventoryResult,
    path: Path | str,
    *,
    dataset_root: str = "",
) -> Path:
    """Write a formatted workbook with Series + Protocol sheets."""
    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError(
            "openpyxl is required for Excel inventory export. "
            "Install with: pip install openpyxl"
        ) from exc

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "DICOM Series Inventory"
    _write_sheet(
        ws1,
        headers=SERIES_INVENTORY_COLUMNS,
        rows=[r.as_ordered() for r in result.series_rows],
    )

    ws2 = wb.create_sheet("Protocol Completeness")
    _write_sheet(
        ws2,
        headers=PROTOCOL_COLUMNS,
        rows=[r.as_ordered() for r in result.protocol_rows],
    )

    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    props = wb.properties
    props.title = "NeuroPipeline DICOM Inventory"
    props.creator = f"NeuroPipeline DICOM Converter {__version__}"
    props.description = (
        f"Dataset: {dataset_root}; generated {generated}; pipeline_version={__version__}"
    )
    props.subject = "DICOM series inventory and protocol completeness"
    props.keywords = f"pipeline_version={__version__}; timestamp={generated}"

    wb.save(out)
    return out


def _write_sheet(ws, *, headers: Sequence[str], rows: list[dict]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    header_font = Font(bold=True, color=HEADER_FONT_COLOR)
    alt_fill = PatternFill("solid", fgColor=ALT_ROW_FILL)
    thin = Border(
        left=Side(style="thin", color="D0D0D0"),
        right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"),
        bottom=Side(style="thin", color="D0D0D0"),
    )

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, header in enumerate(headers, start=1):
            value = row.get(header, "")
            if isinstance(value, (list, tuple)):
                value = "\\".join(str(v) for v in value)
            cell = ws.cell(row=row_idx, column=col_idx, value=_cell_value(value))
            cell.border = thin
            if row_idx % 2 == 0:
                cell.fill = alt_fill

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22

    for col_idx, header in enumerate(headers, start=1):
        letter = get_column_letter(col_idx)
        max_len = len(str(header))
        for row in rows[:200]:  # cap width scan for large inventories
            val = row.get(header, "")
            max_len = max(max_len, len(str(val)) if val is not None else 0)
        ws.column_dimensions[letter].width = min(max(10, max_len + 2), 42)


def _cell_value(value):
    if value is None:
        return ""
    if isinstance(value, float):
        return round(value, 6)
    return value
