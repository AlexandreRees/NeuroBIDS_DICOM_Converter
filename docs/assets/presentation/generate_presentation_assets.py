#!/usr/bin/env python3
"""Generate LinkedIn carousel PNGs + PowerPoint for NeuroPipeline DICOM Converter."""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[3]
LOGO = ROOT / "docs/assets/logo.png"
OUT = ROOT / "docs/assets/presentation"
LI = OUT / "linkedin"
PX = OUT / "pptx_exports"

NAVY = (31, 58, 95)
STEEL = (74, 111, 148)
ACCENT = (46, 139, 140)
CREAM = (247, 246, 243)
INK = (26, 32, 44)
MUTED = (74, 85, 104)
WHITE = (255, 255, 255)
LINE = (226, 232, 240)

CONTENT = [
    {
        "id": "01_hero",
        "kind": "hero",
        "kicker": "Neuroimaging software",
        "title": "DICOM → BIDS\nin one workflow",
        "subtitle": (
            "NeuroPipeline DICOM Converter turns MRI DICOM folders into "
            "analysis-ready NIfTI / BIDS — with preview, inventory, and batch queue."
        ),
        "bullets": ["Built for researchers", "No coding required", "DICOM files stay untouched"],
    },
    {
        "id": "02_problem",
        "kind": "split",
        "kicker": "The problem",
        "title": "Real-world DICOM folders are messy",
        "points": [
            ("Nested exports", "Subjects hidden under visit/raw/MR/DICOM trees"),
            ("Shared PatientIDs", "Anonymized IDs collide across subjects"),
            ("Manual naming", "BIDS entities are easy to get wrong by hand"),
            ("Fragile batches", "Large cohorts need pause, retry, and recovery"),
        ],
    },
    {
        "id": "03_workflow",
        "kind": "flow",
        "kicker": "Simple workflow",
        "title": "From folder to BIDS in minutes",
        "steps": [
            "Select DICOM input",
            "Review analysis",
            "Edit BIDS Preview",
            "Convert or Queue",
        ],
    },
    {
        "id": "04_features",
        "kind": "grid",
        "kicker": "What you get",
        "title": "Built for day-to-day research",
        "cards": [
            ("Recursive discovery", "Finds DICOM at any folder depth — with or without .dcm"),
            ("BIDS Preview", "Edit subject, task, run, datatype before conversion"),
            ("Inventory", "Excel/CSV series tables without writing NIfTI"),
            ("Smart naming rules", "Lab-specific ProtocolName → BIDS entity rules"),
            ("Conversion queue", "Pause, retry, and recover multi-subject jobs"),
            ("Validation + report", "HTML report and provenance after conversion"),
        ],
    },
    {
        "id": "05_safety",
        "kind": "safety",
        "kicker": "Safety first",
        "title": "Source DICOM is never modified",
        "lines": [
            "No rename, move, delete, or overwrite of raw DICOM",
            "Naming edits live only in an in-memory plan / exports",
            "Ideal for lab archives and shared MRI cores",
        ],
    },
    {
        "id": "06_cta",
        "kind": "cta",
        "kicker": "Ready to try",
        "title": "Convert with confidence",
        "subtitle": "Windows desktop app powered by dcm2niix — Convert · Queue · Settings · Logs",
        "chips": ["dcm2niix", "BIDS", "NIfTI", "Inventory", "Queue"],
    },
]


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/liberation/LiberationSans-Bold.ttf"
        if bold
        else "/usr/share/fonts/liberation/LiberationSans-Regular.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
        if bold
        else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for c in candidates:
        if Path(c).exists():
            return ImageFont.truetype(c, size)
    return ImageFont.load_default()


def rounded_rect(draw, box, r, fill, outline=None, width=1):
    draw.rounded_rectangle(box, radius=r, fill=fill, outline=outline, width=width)


def paste_logo(img, xy, size=120):
    logo = Image.open(LOGO).convert("RGBA").resize((size, size), Image.Resampling.LANCZOS)
    img.paste(logo, xy, logo)


def wrap_text(draw, text, font_obj, max_w):
    words = text.split()
    lines, cur = [], ""
    for w in words:
        trial = (cur + " " + w).strip()
        if draw.textlength(trial, font=font_obj) <= max_w:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def draw_footer(draw, w, h, page, total):
    draw.rectangle((0, h - 56, w, h), fill=NAVY)
    f = font(18)
    draw.text(
        (48, h - 38),
        "NeuroPipeline DICOM Converter  ·  v1.0",
        font=f,
        fill=(215, 226, 240),
        anchor="lm",
    )
    draw.text((w - 48, h - 38), f"{page}/{total}", font=f, fill=(215, 226, 240), anchor="rm")


def header_bar(draw, w):
    draw.rectangle((0, 0, w, 10), fill=ACCENT)


def render(c: dict, idx: int, total: int, size: tuple[int, int]) -> Image.Image:
    w, h = size
    wide = w > h
    img = Image.new("RGB", (w, h), CREAM)
    draw = ImageDraw.Draw(img)
    header_bar(draw, w)
    kind = c["kind"]

    if kind == "hero":
        panel = 360 if wide else 220
        title_size = 72 if wide else 64
        left = panel + 60 if wide else 260
        draw.rectangle((0, 10, panel, h - 56), fill=NAVY)
        paste_logo(img, ((panel - (140 if wide else 120)) // 2, 80 if wide else 60), 140 if wide else 120)
        draw.text((48, 90 if wide else 80), c["kicker"].upper(), font=font(24 if wide else 22, True), fill=ACCENT)
        y = 140 if wide else 120
        for line in c["title"].split("\n"):
            draw.text((left, y), line, font=font(title_size, True), fill=NAVY)
            y += 90 if wide else 78
        for ln in wrap_text(draw, c["subtitle"], font(30 if wide else 28), w - left - 80):
            draw.text((left, y + 20), ln, font=font(30 if wide else 28), fill=MUTED)
            y += 42 if wide else 40
        y += 40
        bw = 520 if wide else 420
        for b in c["bullets"]:
            rounded_rect(draw, (left, y, left + bw, y + (58 if wide else 54)), 14, WHITE, LINE, 2)
            draw.ellipse((left + 20, y + 20, left + 42, y + 42), fill=ACCENT)
            draw.text((left + 64, y + 29), b, font=font(26 if wide else 24, True), fill=INK, anchor="lm")
            y += 74 if wide else 70

    elif kind == "split":
        paste_logo(img, (64 if wide else 48, 40 if wide else 36), 80 if wide else 72)
        draw.text(
            (170 if wide else 140, 70 if wide else 58),
            c["kicker"].upper(),
            font=font(22 if wide else 20, True),
            fill=ACCENT,
            anchor="lm",
        )
        draw.text((64 if wide else 48, 150 if wide else 140), c["title"], font=font(52 if wide else 48, True), fill=NAVY)
        y = 250 if wide else 240
        for title, desc in c["points"]:
            box_h = 150 if wide else 140
            rounded_rect(draw, (64 if wide else 48, y, w - (64 if wide else 48), y + box_h), 18, WHITE, LINE, 2)
            draw.rectangle((64 if wide else 48, y, (84 if wide else 64), y + box_h), fill=NAVY)
            draw.text((120 if wide else 96, y + (40 if wide else 36)), title, font=font(32 if wide else 30, True), fill=NAVY)
            draw.text((120 if wide else 96, y + (95 if wide else 82)), desc, font=font(26 if wide else 24), fill=MUTED)
            y += box_h + 20

    elif kind == "flow":
        paste_logo(img, (64 if wide else 48, 40 if wide else 36), 80 if wide else 72)
        draw.text(
            (170 if wide else 140, 70 if wide else 58),
            c["kicker"].upper(),
            font=font(22 if wide else 20, True),
            fill=ACCENT,
            anchor="lm",
        )
        draw.text((64 if wide else 48, 150 if wide else 140), c["title"], font=font(52 if wide else 46, True), fill=NAVY)
        steps = c["steps"]
        box_w = 340 if wide else 200
        gap = 40 if wide else 28
        total_w = len(steps) * box_w + (len(steps) - 1) * gap
        x0 = (w - total_w) // 2
        y0 = 380 if wide else 420
        box_h = 260 if wide else 220
        for i, step in enumerate(steps):
            x = x0 + i * (box_w + gap)
            rounded_rect(draw, (x, y0, x + box_w, y0 + box_h), 22 if wide else 20, WHITE, LINE, 2)
            cx = x + box_w // 2
            r = 40 if wide else 30
            draw.ellipse((cx - r, y0 + 36, cx + r, y0 + 36 + 2 * r), fill=NAVY)
            draw.text((cx, y0 + 36 + r), str(i + 1), font=font(36 if wide else 32, True), fill=WHITE, anchor="mm")
            lines = wrap_text(draw, step, font(28 if wide else 24, True), box_w - 40)
            ty = y0 + (150 if wide else 120)
            for ln in lines:
                draw.text((cx, ty), ln, font=font(28 if wide else 24, True), fill=INK, anchor="ma")
                ty += 36 if wide else 32
            if i < len(steps) - 1:
                draw.polygon(
                    [
                        (x + box_w + 6, y0 + box_h // 2 - 15),
                        (x + box_w + gap - 6, y0 + box_h // 2),
                        (x + box_w + 6, y0 + box_h // 2 + 15),
                    ],
                    fill=STEEL,
                )
        draw.text(
            (w // 2, 780 if wide else 780),
            "Powered by dcm2niix  ·  BIDS-ready NIfTI output",
            font=font(26 if wide else 24),
            fill=MUTED,
            anchor="ma",
        )

    elif kind == "grid":
        paste_logo(img, (64 if wide else 48, 40 if wide else 36), 80 if wide else 72)
        draw.text(
            (170 if wide else 140, 70 if wide else 58),
            c["kicker"].upper(),
            font=font(22 if wide else 20, True),
            fill=ACCENT,
            anchor="lm",
        )
        draw.text((64 if wide else 48, 150 if wide else 140), c["title"], font=font(50 if wide else 44, True), fill=NAVY)
        cards = c["cards"]
        cols = 3 if wide else 2
        card_w = 560 if wide else 460
        card_h = 220 if wide else 200
        gx = 64 if wide else 48
        gy = 260 if wide else 240
        gapx, gapy = (36, 28) if wide else (40, 28)
        for i, (t, d) in enumerate(cards):
            col, row = i % cols, i // cols
            x = gx + col * (card_w + gapx)
            y = gy + row * (card_h + gapy)
            rounded_rect(draw, (x, y, x + card_w, y + card_h), 18, WHITE, LINE, 2)
            draw.rectangle((x, y, x + 12, y + card_h), fill=ACCENT if i % 2 == 0 else NAVY)
            draw.text((x + 40, y + 40), t, font=font(28 if wide else 26, True), fill=NAVY)
            ty = y + 100 if wide else y + 90
            for ln in wrap_text(draw, d, font(24 if wide else 22), card_w - 70):
                draw.text((x + 40, ty), ln, font=font(24 if wide else 22), fill=MUTED)
                ty += 32 if wide else 30

    elif kind == "safety":
        draw.rectangle((0, 10, w, h - 56), fill=NAVY)
        paste_logo(img, (w // 2 - (70 if wide else 60), 70 if wide else 80), 140 if wide else 120)
        draw.text((w // 2, 250 if wide else 240), c["kicker"].upper(), font=font(24 if wide else 22, True), fill=ACCENT, anchor="ma")
        draw.text((w // 2, 320 if wide else 300), c["title"], font=font(54 if wide else 48, True), fill=WHITE, anchor="ma")
        y = 430 if wide else 420
        for line in c["lines"]:
            left = 280 if wide else 100
            rounded_rect(draw, (left, y, w - left, y + (100 if wide else 90)), 18 if wide else 16, (36, 66, 105))
            draw.text((left + 60 if wide else 140, y + (50 if wide else 45)), "✓  " + line, font=font(28 if wide else 26), fill=WHITE, anchor="lm")
            y += 120 if wide else 110

    elif kind == "cta":
        paste_logo(img, (64 if wide else 48, 40 if wide else 36), 80 if wide else 72)
        draw.text(
            (170 if wide else 140, 70 if wide else 58),
            c["kicker"].upper(),
            font=font(22 if wide else 20, True),
            fill=ACCENT,
            anchor="lm",
        )
        draw.text((64 if wide else 48, 200), c["title"], font=font(64 if wide else 56, True), fill=NAVY)
        y = 320 if wide else 300
        for ln in wrap_text(draw, c["subtitle"], font(32 if wide else 28), w - 160):
            draw.text((64 if wide else 48, y), ln, font=font(32 if wide else 28), fill=MUTED)
            y += 44 if wide else 40
        x = 64 if wide else 48
        y = 480 if wide else 460
        for chip in c["chips"]:
            tw = int(draw.textlength(chip, font=font(26 if wide else 24, True))) + (56 if wide else 48)
            rounded_rect(draw, (x, y, x + tw, y + (60 if wide else 56)), 30 if wide else 28, NAVY)
            draw.text((x + tw // 2, y + (30 if wide else 28)), chip, font=font(26 if wide else 24, True), fill=WHITE, anchor="mm")
            x += tw + 18
            if x > w - 220:
                x = 64 if wide else 48
                y += 76
        rounded_rect(draw, (64 if wide else 48, 680 if wide else 700, w - (64 if wide else 48), 860 if wide else 820), 22 if wide else 20, WHITE, LINE, 2)
        draw.text(
            (w // 2, 740 if wide else 740),
            "Windows desktop  ·  Research-ready  ·  BIDS-focused",
            font=font(30 if wide else 26, True),
            fill=NAVY,
            anchor="ma",
        )
        draw.text(
            (w // 2, 800 if wide else 785),
            "Convert  ·  Queue  ·  Settings  ·  Logs",
            font=font(28 if wide else 24),
            fill=MUTED,
            anchor="ma",
        )

    if kind != "safety":
        draw_footer(draw, w, h, idx, total)
    else:
        draw.text((w - (64 if wide else 48), h - 38), f"{idx}/{total}", font=font(18), fill=(215, 226, 240), anchor="rm")
        draw.text(
            (64 if wide else 48, h - 38),
            "NeuroPipeline DICOM Converter  ·  v1.0",
            font=font(18),
            fill=(215, 226, 240),
            anchor="lm",
        )
    return img


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    LI.mkdir(exist_ok=True)
    PX.mkdir(exist_ok=True)
    n = len(CONTENT)
    meta = []
    wide_paths = []
    for i, c in enumerate(CONTENT, 1):
        sq = render(c, i, n, (1080, 1080))
        p = LI / f"linkedin_{c['id']}.png"
        sq.save(p, "PNG", optimize=True)
        meta.append((p.name, c["title"].replace("\n", " ")))
        print("wrote", p)

        wide = render(c, i, n, (1920, 1080))
        wp = PX / f"slide_{c['id']}.png"
        wide.save(wp, "PNG", optimize=True)
        wide_paths.append(wp)
        print("wrote", wp)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for p in wide_paths:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(p), Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height)
    ppt = OUT / "NeuroPipeline_DICOM_Converter_overview.pptx"
    prs.save(ppt)
    print("wrote", ppt)

    lines = [
        "# NeuroPipeline presentation assets\n\n",
        "English pack for LinkedIn + PowerPoint.\n\n",
        "## LinkedIn carousel (1080×1080)\n",
        "Upload **in this order** as a LinkedIn document/carousel:\n\n",
    ]
    for i, (name, title) in enumerate(meta, 1):
        lines.append(f"{i}. `linkedin/{name}` — {title}\n")
    lines += [
        "\n## PowerPoint\n",
        f"- `{ppt.name}` — 6 widescreen slides (16:9)\n",
        "- PNG sources also in `pptx_exports/`\n\n",
        "## Suggested LinkedIn caption\n\n",
        "> New tool for neuroimaging teams: **NeuroPipeline DICOM Converter** turns messy MRI DICOM folders into BIDS-ready NIfTI — with recursive discovery, editable BIDS preview, inventory Excel export, naming rules, and a conversion queue.\n>\n",
        "> Source DICOM is never modified.\n>\n",
        "> #Neuroimaging #BIDS #MRI #OpenScience #ResearchSoftware\n",
    ]
    (OUT / "README.md").write_text("".join(lines), encoding="utf-8")


if __name__ == "__main__":
    main()
