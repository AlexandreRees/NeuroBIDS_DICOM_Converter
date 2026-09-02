"""Unit tests for DICOM inventory (no real DICOM tree required)."""

from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest

from neuro_pipeline.inventory.csv_exporter import export_inventory_csv
from neuro_pipeline.inventory.manager import InventoryManager
from neuro_pipeline.inventory.models import (
    PROTOCOL_ACQUISITIONS,
    SERIES_INVENTORY_COLUMNS,
    InventoryResult,
    SeriesInventoryRow,
)
from neuro_pipeline.inventory.protocol_completeness import (
    build_protocol_rows,
    detect_acquisition_keys,
)
from neuro_pipeline.models import DicomSeries, SeriesStatus


def _series(
    description: str,
    *,
    seq: str = "anat",
    fine: str = "ANAT_T1",
    patient_id: str = "SUB001",
    series_number: int = 1,
    series_uid: str = "",
) -> DicomSeries:
    root = Path("fake_dicom")
    return DicomSeries(
        patient_id=patient_id,
        study_description="Study",
        series_description=description,
        protocol_name=description,
        series_number=series_number,
        acquisition_number=1,
        modality="MR",
        num_images=32,
        source_dir=root,
        sample_file=root / "x.dcm",
        status=SeriesStatus.PENDING,
        sequence_type=seq,
        fine_sequence_type=fine,
        smart_name=description,
        series_instance_uid=series_uid or f"1.2.3.{series_number}",
        study_instance_uid="1.2.3",
    )


def test_detect_acquisition_keys_core() -> None:
    assert "T1w" in detect_acquisition_keys(_series("t1_mprage", fine="ANAT_T1"))
    assert "FLAIR" in detect_acquisition_keys(
        _series("t2_flair", seq="anat", fine="FLAIR")
    )
    assert "REST_AP" in detect_acquisition_keys(
        _series("rest_AP", seq="func", fine="FMRI_REST")
    )
    assert "FieldMap_PA" in detect_acquisition_keys(
        _series("fmap_PA", seq="fmap", fine="FMAP")
    )
    assert "DWI" in detect_acquisition_keys(_series("ep2d_diff", seq="dwi", fine="DWI"))
    assert "SBRef" in detect_acquisition_keys(
        _series("bold_sbref", seq="func", fine="FMRI_REST")
    )


def test_protocol_completeness_movies_and_complete_flag() -> None:
    series = [
        _series("t1_mprage", fine="ANAT_T1", series_number=1),
        _series("t2_flair", fine="FLAIR", series_number=2),
        _series("rest_AP", seq="func", fine="FMRI_REST", series_number=3),
        _series("rest_PA", seq="func", fine="FMRI_REST", series_number=4),
        _series("movie_run1", seq="func", fine="FMRI_TASK", series_number=5),
        _series("movie_run2", seq="func", fine="FMRI_TASK", series_number=6),
        _series("movie_run3", seq="func", fine="FMRI_TASK", series_number=7),
        _series("movie_run4", seq="func", fine="FMRI_TASK", series_number=8),
        _series("dwi", seq="dwi", fine="DWI", series_number=9),
        _series("bold_sbref", seq="func", fine="FMRI_REST", series_number=10),
        _series("fmap_AP", seq="fmap", fine="FMAP", series_number=11),
        _series("fmap_PA", seq="fmap", fine="FMAP", series_number=12),
    ]
    rows = build_protocol_rows(series, session_override="01")
    assert len(rows) == 1
    row = rows[0]
    ordered = row.as_ordered()
    assert ordered["Subject"]
    assert ordered["Session"] == "01"
    for key in PROTOCOL_ACQUISITIONS:
        assert ordered[key] == "✓", key
    assert ordered["Complete Protocol"] == "TRUE"


def test_protocol_incomplete_missing_flair() -> None:
    rows = build_protocol_rows(
        [_series("t1_mprage", fine="ANAT_T1")],
        subject_override="P1",
        session_override="01",
    )
    assert rows[0].as_ordered()["FLAIR"] == "✗"
    assert rows[0].as_ordered()["Complete Protocol"] == "FALSE"


def test_inventory_build_plans_bids_names_without_dicom_io(tmp_path: Path) -> None:
    series = [
        _series("t1_mprage", fine="ANAT_T1", series_number=1),
        _series("t1_mprage_repeat", fine="ANAT_T1", series_number=2),
        _series("rest_AP", seq="func", fine="FMRI_REST", series_number=3),
    ]
    empty_meta = MagicMock()
    empty_meta.study_instance_uid = ""
    empty_meta.series_instance_uid = ""
    empty_meta.series_number = None
    empty_meta.series_description = ""
    empty_meta.protocol_name = ""
    empty_meta.modality = "MR"
    empty_meta.num_images = 0
    empty_meta.image_type = []
    empty_meta.manufacturer = ""
    empty_meta.manufacturer_model_name = ""
    empty_meta.magnetic_field_strength = None
    empty_meta.rows = None
    empty_meta.columns = None
    empty_meta.pixel_spacing = None
    empty_meta.slice_thickness = None
    empty_meta.echo_time = None
    empty_meta.repetition_time = None
    empty_meta.flip_angle = None
    empty_meta.model_dump.return_value = {}

    extractor = MagicMock()
    extractor.extract_from_file.return_value = empty_meta

    manager = InventoryManager(metadata_extractor=extractor)
    with patch(
        "neuro_pipeline.inventory.manager.read_optional_tags",
        return_value={},
    ):
        result = manager.build(
            series,
            dataset_root=tmp_path,
            subject_override="01",
            session_override="01",
        )

    assert result.n_series == 3
    assert result.n_subjects == 1
    rows = [r.as_ordered() for r in result.series_rows]
    assert list(rows[0].keys()) == list(SERIES_INVENTORY_COLUMNS)
    assert rows[0]["Subject"] == "01"
    assert rows[0]["Session"] == "01"
    assert rows[0]["Expected BIDS Datatype"] == "anat"
    assert rows[0]["Expected BIDS Suffix"] == "T1w"
    assert rows[0]["Planned Filename"].endswith("_T1w.nii.gz")
    # Second T1 gets a run entity via unique_stem
    assert rows[1]["Planned Run Number"] in {"1", "01", "1".zfill(2)} or "run-" in rows[
        1
    ]["Planned Filename"]
    assert rows[2]["Expected BIDS Datatype"] == "func"
    assert rows[2]["Conversion Status"] == "Pending"
    assert result.protocol_rows[0].as_ordered()["T1w"] == "✓"


def test_export_xlsx_and_csv(tmp_path: Path) -> None:
    pytest.importorskip("openpyxl")
    from neuro_pipeline.inventory.excel_exporter import export_inventory_xlsx
    from openpyxl import load_workbook

    result = InventoryResult(
        series_rows=[
            SeriesInventoryRow(
                values={
                    "Dataset Root": str(tmp_path),
                    "Subject": "01",
                    "Session": "01",
                    "SeriesDescription": "t1",
                    "Detected Acquisition Type": "ANAT_T1",
                    "Expected BIDS Datatype": "anat",
                    "Expected BIDS Suffix": "T1w",
                    "Planned Filename": "sub-01_ses-01_T1w.nii.gz",
                    "Conversion Status": "Pending",
                }
            )
        ],
        protocol_rows=build_protocol_rows(
            [_series("t1_mprage", fine="ANAT_T1")],
            subject_override="01",
            session_override="01",
        ),
        n_series=1,
        n_subjects=1,
    )
    xlsx = export_inventory_xlsx(result, tmp_path / "dicom_inventory.xlsx", dataset_root=str(tmp_path))
    assert xlsx.is_file()
    wb = load_workbook(xlsx)
    assert "DICOM Series Inventory" in wb.sheetnames
    assert "Protocol Completeness" in wb.sheetnames
    ws = wb["DICOM Series Inventory"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref
    assert ws["A1"].value == "Dataset Root"
    assert wb.properties.creator and "NeuroPipeline" in wb.properties.creator

    csv_path = export_inventory_csv(result, tmp_path / "dicom_inventory.csv")
    assert csv_path.is_file()
    text = csv_path.read_text(encoding="utf-8")
    assert "SeriesDescription" in text
    assert (tmp_path / "dicom_inventory_protocol_completeness.csv").is_file()


def test_manager_run_exports(tmp_path: Path) -> None:
    pytest.importorskip("openpyxl")
    series = [_series("t1_mprage", fine="ANAT_T1")]
    empty_meta = MagicMock()
    for attr in (
        "study_instance_uid",
        "series_instance_uid",
        "series_description",
        "protocol_name",
        "modality",
        "manufacturer",
        "manufacturer_model_name",
    ):
        setattr(empty_meta, attr, "")
    for attr in (
        "series_number",
        "rows",
        "columns",
        "magnetic_field_strength",
        "pixel_spacing",
        "slice_thickness",
        "echo_time",
        "repetition_time",
        "flip_angle",
    ):
        setattr(empty_meta, attr, None)
    empty_meta.num_images = 0
    empty_meta.image_type = []
    empty_meta.model_dump.return_value = {}
    extractor = MagicMock()
    extractor.extract_from_file.return_value = empty_meta
    manager = InventoryManager(metadata_extractor=extractor)
    with patch(
        "neuro_pipeline.inventory.manager.read_optional_tags",
        return_value={},
    ):
        result = manager.run(
            tmp_path / "in",
            tmp_path / "out",
            series_list=series,
            subject_override="01",
            session_override="01",
            write_csv=True,
        )
    assert Path(result.xlsx_path).is_file()
    assert Path(result.csv_path).is_file()
